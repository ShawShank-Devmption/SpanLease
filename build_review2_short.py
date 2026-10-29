#!/usr/bin/env python3
"""Generate the concise SpanLease Project Review 2 presentation."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = Path("output/SpanLease_Review2_Presentation.pptx")
ASSETS = Path(__file__).resolve().parent / "output" / ".review2_assets"
P = Presentation()
P.slide_width, P.slide_height = Inches(13.333), Inches(7.5)
BLANK = P.slide_layouts[6]

BG = RGBColor(248, 246, 243)
INK = RGBColor(37, 42, 48)
MUTED = RGBColor(83, 91, 99)
TERRA = RGBColor(176, 74, 48)
NAVY = RGBColor(39, 56, 78)
GREEN = RGBColor(52, 111, 82)
LINE = RGBColor(215, 208, 200)
PALE = RGBColor(239, 233, 226)
WHITE = RGBColor(255, 255, 255)


def shape(s, x, y, w, h, fill=WHITE, border=None, kind=MSO_SHAPE.RECTANGLE):
    q = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    q.fill.solid()
    q.fill.fore_color.rgb = fill
    if border:
        q.line.color.rgb = border
        q.line.width = Pt(1.2)
    else:
        q.line.fill.background()
    return q


def txt(s, x, y, w, h, text, size=20, color=INK, bold=False,
        font="Aptos", align=PP_ALIGN.LEFT, margin=0, valign=MSO_ANCHOR.MIDDLE):
    q = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    f = q.text_frame
    f.clear()
    f.word_wrap = True
    f.margin_left = f.margin_right = Inches(margin)
    f.margin_top = f.margin_bottom = 0
    f.vertical_anchor = valign
    para = f.paragraphs[0]
    para.alignment = align
    para.space_before = para.space_after = Pt(0)
    run = para.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return q


def base(title, section, number):
    s = P.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    shape(s, 0, 0, 0.19, 7.5, TERRA)
    txt(s, .65, .42, 10.8, .3, section.upper(), 11, TERRA, True)
    txt(s, .65, .83, 12, .66, title, 32, INK, True, font="Georgia")
    shape(s, .67, 1.63, 1.0, .045, TERRA)
    txt(s, .67, 7.08, 10.4, .22, "SpanLease  ·  Project Review 2", 9, MUTED)
    txt(s, 12.05, 7.08, .55, .22, f"{number:02}", 9, MUTED, align=PP_ALIGN.RIGHT)
    return s


def note(s, text):
    s.notes_slide.notes_text_frame.text = text


def line(s, x, y, w):
    shape(s, x, y, w, .015, LINE)


def arrow(s, x, y, w=.28):
    shape(s, x, y, w, .25, TERRA, kind=MSO_SHAPE.RIGHT_ARROW)


# 1 — title
s = P.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = BG
shape(s, 0, 0, .22, 7.5, TERRA)
txt(s, .82, .77, 10, .33, "BCSE303P OPERATING SYSTEMS LAB · PROJECT REVIEW 2", 13, TERRA, True)
txt(s, .79, 1.52, 11.5, 1.12, "SpanLease", 57, INK, True, font="Georgia")
txt(s, .83, 2.82, 10.9, 1.1,
    "Detection of persistent circular waits across blocking RPCs and bounded application workers",
    24, NAVY, font="Georgia")
shape(s, .83, 4.67, 11.58, .017, LINE)
txt(s, .83, 5.0, 11.2, .57,
    "Architecture  ·  modules  ·  results  ·  output screenshots  ·  live demo", 18, MUTED)
txt(s, .83, 6.56, 10.4, .36, "Shashank  ·  B.Tech CSE (Data Science), VIT", 13, MUTED)
note(s, "Rubric: System Design and Architecture 1; Module Identification 1; Module Completion 4; Presentation and Discussion 4. Source: supplied BCSE303P Project Review Rubrics PDF.")

# 2 — problem and scope
s = base("The operating-systems problem", "Problem", 2)
txt(s, .72, 2.0, 11.8, .75,
    "A handler holds a worker slot while making a synchronous downstream call.", 25, NAVY, True)
txt(s, .72, 2.9, 11.7, .66,
    "A cycle becomes stuck when every service it reaches also needs a slot that the cycle already holds.",
    22)
shape(s, .75, 4.11, 11.78, 1.37, PALE)
txt(s, 1.03, 4.28, 10.9, .35, "SUPPORTED MODEL", 12, TERRA, True)
txt(s, 1.03, 4.72, 10.9, .54,
    "Blocking unary Java/gRPC calls · explicit bounded application pools · one outstanding downstream call per execution",
    18, INK)
txt(s, .76, 6.04, 11.75, .48,
    "Goal: localize a persistent closed wait before the relevant RPC deadline when complete evidence permits it.",
    18, GREEN, True)
note(s, "Sources: requirements.md §§2–3 and design.md §§2–3. This is a research objective, not a measured timing result.")

# 3 — paper-style system architecture
s = base("System architecture", "Design and architecture", 3)
txt(s, .7, 1.92, 12, .35,
    "A layered monitoring pipeline with the application, observation path, and analysis separated",
    16, MUTED)
labels = [
    ("APPLICATION", .75, 2.46, 2.15, NAVY),
    ("OBSERVATION", .75, 4.02, 2.15, TERRA),
    ("ANALYSIS", .75, 5.47, 2.15, GREEN),
]
for label, x, y, w, c in labels:
    shape(s, x, y, w, .42, c)
    txt(s, x+.11, y+.03, w-.2, .35, label, 12, WHITE, True)

for x, w, head, body in [
    (3.04, 2.56, "gRPC arrival", "Early stream hook and request identity"),
    (5.96, 2.57, "Application gate", "Bounded queue, named slot ownership"),
    (8.88, 3.52, "Blocking helper", "Call interval, causal request and response refs"),
]:
    shape(s, x, 2.42, w, 1.14, WHITE, LINE)
    txt(s, x+.13, 2.51, w-.25, .36, head, 17, NAVY, True)
    txt(s, x+.13, 2.94, w-.25, .47, body, 12, MUTED)
arrow(s, 5.66, 2.88)
arrow(s, 8.59, 2.88)

for x, w, head, body in [
    (3.04, 3.28, "Custom OTel LogRecords", "Ordered transitions, leases and checkpoints"),
    (6.75, 2.43, "OTLP transport", "Bounded async export"),
    (9.62, 2.78, "Receiver", "Validate and admit events"),
]:
    shape(s, x, 4.0, w, 1.07, WHITE, LINE)
    txt(s, x+.13, 4.08, w-.25, .35, head, 17, TERRA, True)
    txt(s, x+.13, 4.49, w-.25, .41, body, 12, MUTED)
arrow(s, 6.39, 4.38)
arrow(s, 9.25, 4.38)

shape(s, 3.04, 5.48, 5.18, .97, WHITE, LINE)
txt(s, 3.18, 5.55, 4.9, .35, "Analyzer", 17, GREEN, True)
txt(s, 3.18, 5.94, 4.9, .38, "Certified cut, graph reduction, persistence", 12, MUTED)
arrow(s, 8.34, 5.82)
shape(s, 8.72, 5.48, 3.68, .97, WHITE, LINE)
txt(s, 8.86, 5.55, 3.4, .35, "Evidence report", 17, GREEN, True)
txt(s, 8.86, 5.94, 3.4, .38, "Verdict, scope, cut and coverage", 12, MUTED)
note(s, "Architecture adapted in figure style from the system/dataflow presentation of DDMon (Rowicki et al., OOPSLA 2025), Hindsight (Zhang et al., NSDI 2023), and CRISP (Zhang et al., USENIX ATC 2022). Technical architecture follows design.md §§2, 4–6. OTLP receiver and full live wiring remain integration work. The analyzer never reads independent ground truth.")

# 4 — modules and completion
s = base("Modules and current completion", "Module identification and completion", 4)
rows = [
    ("Common contract", "v2 event, manifest and report DTOs with validation", "Implemented"),
    ("Runtime instrumentation", "Slot gate, blocking-call lifecycle and bounded OTel sink", "Component tests pass"),
    ("Analyzer", "OTLP decoding, certified evidence and conservative verdict core", "Deterministic demo passes"),
    ("Testbed", "Unary gRPC contract and real loopback lifecycle feasibility", "Feasibility tested"),
]
y = 2.04
for i, (name, desc, status) in enumerate(rows):
    line(s, .75, y+1.04, 11.75)
    txt(s, .78, y+.05, 2.95, .38, name, 18, NAVY, True)
    txt(s, 3.76, y+.03, 6.52, .78, desc, 17)
    txt(s, 10.34, y+.07, 2.03, .58, status, 13, GREEN, True, align=PP_ALIGN.RIGHT)
    y += 1.11
txt(s, .78, 6.66, 11.6, .28,
    "Current evidence is component-level. Real-time, multi-service integration is the next milestone.",
    13, TERRA, True)
note(s, "Sources: common/, instrumentation/, analyzer/, testbed/ source and targeted tests; docs/mvp-demo-guide.md. Do not describe component tests as end-to-end detector validation. The current G0 record is solo role-based review under D13/D14, not independent three-person approval.")

# 5 — detection algorithm
s = base("Detection algorithm: certified wait-for reduction", "Algorithm", 5)
steps = [
    ("1", "Certify observations", "Require init and contiguous sequence prefixes through checkpoints."),
    ("2", "Construct a causal cut", "Join exact RPC parents; reconstruct queue, ownership and blocked calls."),
    ("3", "Reduce progress paths", "Remove waits served by a free slot or an owner/counterpart able to progress."),
    ("4", "Check persistence", "The same closed dependencies must overlap continuously for at least τ."),
    ("5", "Report evidence", "Return deadlock, no deadlock, inconclusive or insufficient observability."),
]
y = 2.01
for n, head, body in steps:
    shape(s, .77, y+.035, .46, .46, TERRA)
    txt(s, .77, y+.053, .46, .38, n, 16, WHITE, True, align=PP_ALIGN.CENTER)
    txt(s, 1.43, y, 3.42, .52, head, 19, NAVY, True)
    txt(s, 4.8, y-.025, 7.55, .75, body, 17)
    y += .89
shape(s, .78, 6.53, 11.61, .025, LINE)
txt(s, .78, 6.63, 11.56, .24,
    "A missing record or unknown slot state prevents a confirmed verdict for that scope.", 13, TERRA, True)
note(s, "Algorithm source: design.md §5.1–5.6. The worklist starts with observed blocked executions and queued requests, removes progress-capable members to a fixed point, and temporally qualifies a residual. The temporal pruning step cannot produce a negative verdict. The analyzer implementation and deterministic demos are in analyzer/src/main/java. Formal proof and end-to-end evaluation are not claimed.")

# 6 — synthetic workload and consulted research
s = base("Synthetic RPC workload and source basis", "Dataset", 6)
txt(s, .76, 1.95, 11.7, .58,
    "We use deterministic synthetic scenarios that model synchronous service-to-service RPC chains.",
    21, NAVY, True)
shape(s, .78, 2.83, 11.58, 1.61, PALE)
txt(s, 1.06, 3.01, 5.22, .35, "WORKLOAD SHAPE", 12, TERRA, True)
txt(s, 1.06, 3.43, 5.04, .82,
    "CallRequest carries a request ID and remaining hops; each hop selects a service and work duration.",
    17)
txt(s, 6.48, 3.01, 5.31, .35, "DETECTOR INPUT", 12, TERRA, True)
txt(s, 6.48, 3.43, 5.2, .82,
    "v2 JSON event envelopes represent RPC causality, queue waits, slot ownership and time evidence.",
    17)
txt(s, .78, 4.87, 11.55, .41,
    "Cases exercised: closed A–B–C wait, spare capacity, missing event, expired lease and short persistence.",
    16, INK)
line(s, .78, 5.47, 11.55)
txt(s, .78, 5.65, 2.13, .39, "Consulted", 17, NAVY, True)
txt(s, 2.95, 5.61, 9.42, .83,
    "DDMon for distributed monitor semantics; Hindsight and CRISP for telemetry and diagnosis context; Palette and a microservice graph generator for future trace-derived workload design.",
    15)
txt(s, .78, 6.62, 11.56, .29,
    "Real gRPC loopback behavior is tested; the full live multi-service dataset is planned for integration.",
    12, TERRA, True)
note(s, "Sources: testbed/src/main/proto/spanlease/testbed/chain.proto; analyzer/src/main/java/dev/spanlease/analyzer/ProductionDemo.java; docs/scenario-truth-spec.md; docs/literature-review.md. gRPC uses protobuf on the wire. JSON here describes SpanLease telemetry and reports. The current deterministic analyzer replay uses constructed observations, while a separate feasibility test performs a real loopback gRPC call. No production or third-party trace corpus has been ingested. Palette: Anand et al., arXiv:2506.06448. Microservice graph generator: Du et al., arXiv:2412.19083. Neither was used to generate the current scenarios.")

# 7 — scenario results and analysis
s = base("Deterministic results and analysis", "Results and analysis", 7)
shape(s, .77, 2.06, 5.53, 2.43, WHITE, LINE)
txt(s, 1.05, 2.32, 4.9, .5, "Closed three-service wait", 20, NAVY, True)
txt(s, 1.05, 2.98, 4.94, .54, "CONFIRMED_DEADLOCK", 22, TERRA, True)
txt(s, 1.05, 3.7, 4.96, .51,
    "All required workers are held by the closed dependency set.", 16)
shape(s, 6.68, 2.06, 5.53, 2.43, WHITE, LINE)
txt(s, 6.96, 2.32, 4.94, .5, "Same chain, one spare slot", 20, NAVY, True)
txt(s, 6.96, 2.98, 4.94, .54, "CONFIRMED_NO_DEADLOCK", 20, GREEN, True)
txt(s, 6.96, 3.7, 4.96, .51,
    "The available capacity breaks the closed wait at the certified cut.", 16)
txt(s, .79, 5.01, 11.42, .69,
    "The closed ring persists for 2,670 ms against τ = 2,000 ms. One additional free slot removes the residual.",
    18, INK, True)
shape(s, .78, 5.95, 11.45, .035, TERRA)
txt(s, .78, 6.17, 11.42, .51,
    "Next: real-time integration from live gRPC through OTLP ingestion to the analyzer.",
    20, TERRA, True)
note(s, "Source: freshly run analyzer/src/main/java/dev/spanlease/analyzer/ProductionDemo.java and its v2 JSON reports. The closed-ring report has 3/3 observed occupied units, seven evidence edges, a 2670 ms narrowed interval, and tau 2000 ms. The spare-capacity report has 4/4 observed units and an empty residual. These are deterministic production-analyzer demonstrations, not a live multi-service end-to-end experiment.")


def output_slide(number, title, path, caption, crop_top, crop_bottom,
                 crop_right, x, y, width, height, source_note):
    slide = base(title, "Screenshots of outputs", number)
    picture = slide.shapes.add_picture(str(ASSETS / path), Inches(x), Inches(y),
                                       width=Inches(width), height=Inches(height))
    picture.crop_top = crop_top
    picture.crop_bottom = crop_bottom
    picture.crop_right = crop_right
    txt(slide, .85, 6.80, 11.53, .20, caption, 10, MUTED)
    note(slide, source_note)
    return slide


# 8–10 — genuine captures of current outputs, with PowerPoint crop for legibility
output_slide(
    8, "Output 1: closed wait", "deadlock.png",
    "Captured production-analyzer output, selected fields from the v2 JSON report",
    75/768, 248/768, 270/1310, .86, 1.86, 11.45, 4.90,
    "Screenshot of locally rendered fields extracted from the exact output of ./gradlew -q :analyzer:productionDemo under Java 21. Full unedited output: output/SpanLease_Review2_Demo_Output.txt. The browser screenshot is cropped in the slide; data values are unchanged."
)
output_slide(
    9, "Output 2: spare capacity", "spare.png",
    "Captured production-analyzer output, selected fields from the v2 JSON report",
    75/768, 248/768, 270/1310, .86, 1.86, 11.45, 4.90,
    "Screenshot of locally rendered fields extracted from the exact output of ./gradlew -q :analyzer:productionDemo under Java 21. Full unedited output: output/SpanLease_Review2_Demo_Output.txt. The browser screenshot is cropped in the slide; data values are unchanged."
)
output_slide(
    10, "Output 3: verification", "tests.png",
    "Captured summary of fresh Gradle JUnit XML results, 29 tests and zero failures",
    75/768, 223/768, 270/1310, 1.05, 1.84, 11.15, 5.04,
    "Screenshot of locally rendered counts from fresh Gradle JUnit XML results after ./gradlew :instrumentation:test :analyzer:test :testbed:test --offline --rerun-tasks on Java 21. Six instrumentation, nineteen analyzer and four testbed tests passed, with zero failures and errors. The browser screenshot is cropped in the slide."
)

# 11 — live demonstration setup
s = base("Live demonstration", "Presenter sequence", 11)
txt(s, .77, 2.00, 11.63, .48,
    "Run from the project directory with Java 21:", 21, NAVY, True)
shape(s, .77, 2.70, 11.59, 1.00, NAVY)
txt(s, 1.03, 2.94, 11.06, .47,
    "bash output/SpanLease_Review2_Live_Demo.sh", 21, WHITE, True, font="Consolas")
steps = [
    ("1", "Real gRPC feasibility", "The loopback test shows arrival, bounded gate admission and blocking-call lifecycle."),
    ("2", "Closed-ring result", "Point to the occupied slots, evidence edges and 2,670 ms persistent wait."),
    ("3", "Spare-capacity control", "Show how the extra free slot reduces the candidate graph to empty."),
]
y = 4.00
for n, head, body in steps:
    txt(s, .80, y, .39, .37, n, 18, TERRA, True)
    txt(s, 1.28, y-.04, 3.15, .43, head, 18, NAVY, True)
    txt(s, 4.49, y-.06, 7.80, .61, body, 16)
    y += .79
txt(s, .80, 6.52, 11.50, .32,
    "Live multi-service OTLP-to-analyzer integration remains the next milestone.",
    15, TERRA, True)
note(s, "Use output/SpanLease_Review2_Live_Demo.sh. It selects Java 21, runs the real loopback gRPC feasibility test, runs the production analyzer demonstration, summarizes both reports, and exits nonzero if either verdict differs. The demo is a component demonstration, not a real-time end-to-end detector run.")

OUT.parent.mkdir(parents=True, exist_ok=True)
P.save(OUT)
print(f"{OUT} ({len(P.slides)} slides)")
